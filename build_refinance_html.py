#!/usr/bin/env python3
"""build_refinance_html.py — 從 grand_pivot_deck.pptx 抽 11 頁內容 → 深色手機友善 HTML。
用法：python build_refinance_html.py（需先跑 build_final.py 產出 pptx）"""
import os
import sys
from pptx import Presentation
from pptx.util import Emu

BASE = os.path.dirname(os.path.abspath(__file__))
PPTX = os.path.join(BASE, 'grand_pivot_deck.pptx')
OUT = os.path.join(BASE, 'grand_pivot_deck.html')

prs = Presentation(PPTX)

def shape_top_in(sh):
    try:
        return Emu(sh.top).inches if sh.top is not None else 999
    except Exception:
        return 999

slides_html = []
for idx, slide in enumerate(prs.slides, 1):
    shapes = [sh for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
    shapes.sort(key=shape_top_in)
    title, subtitle, body = '', '', []
    for sh in shapes:
        txt = sh.text_frame.text.strip()
        if not txt:
            continue
        top = shape_top_in(sh)
        if top < 1.0:          # 標題（top=0.3）
            title = txt
        elif top < 1.6:        # 副標（top=1.2）
            subtitle = txt
        else:                  # 內容（top=1.9+）
            body.append(txt)
    body_lines = []
    for b in body:
        for ln in b.split('\n'):
            s = ln.strip()
            body_lines.append(s)
    # 過濾純空行（保留單一空行作為段落間距）
    kept = []
    prev_blank = False
    for s in body_lines:
        blank = (s == '')
        if blank and prev_blank:
            continue
        kept.append(s)
        prev_blank = blank
    while kept and kept[-1] == '':
        kept.pop()

    paras = []
    for s in kept:
        if s == '':
            paras.append('<div class="gap"></div>')
        elif s.startswith(('1️⃣','2️⃣','3️⃣','4️⃣','5️⃣')):
            paras.append(f'<p class="act"><span class="act-num">{s[:2]}</span><span>{s[2:].strip()}</span></p>')
        elif s.startswith('🚫'):
            paras.append(f'<p class="no">{s}</p>')
        elif s.startswith(('📉','📈','📊','💵','📡','🔵','💡','⚠️','📋','🔴','🟡','🟢','🏠','💰','🎯','🔥','📅','✅')):
            paras.append(f'<p class="item">{s}</p>')
        else:
            paras.append(f'<p class="sub-item">{s}</p>')

    slides_html.append(f'''
    <section class="slide" id="s{idx}">
        <div class="slide-head">
            <span class="page-num">{idx:02d} / 11</span>
            <h2>{title}</h2>
            {f'<p class="subtitle">{subtitle}</p>' if subtitle else ''}
        </div>
        <div class="slide-body">{''.join(paras)}</div>
    </section>''')

html = f'''<!DOCTYPE html>
<html lang="zh-Hant">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>大轉向資產配置策略（2026-09-06）</title>
<script src="assets/tailwind.min.js"></script>
<style>
    body {{ background:#0B0D1A; color:#E2E8F0; font-family:system-ui,-apple-system,"Segoe UI","Microsoft JhengHei",sans-serif; }}
    .wrap {{ max-width:880px; margin:0 auto; padding:16px; }}
    header {{ position:sticky; top:0; z-index:10; background:rgba(11,13,26,.92); backdrop-filter:blur(6px);
             border-bottom:1px solid #1E293B; padding:10px 0; }}
    h1 {{ font-size:20px; font-weight:800; color:#fff; margin:0; }}
    .meta {{ font-size:12px; color:#8A8FA0; }}
    .slide {{ background:#111527; border:1px solid #1E293B; border-radius:14px; margin:14px 0; padding:18px; }}
    .slide-head {{ border-bottom:1px solid #232B40; padding-bottom:10px; margin-bottom:12px; }}
    .page-num {{ font-size:11px; color:#F7A01C; font-weight:700; letter-spacing:.08em; }}
    h2 {{ font-size:22px; font-weight:800; color:#fff; margin:4px 0 0; }}
    .subtitle {{ font-size:13px; color:#8A8FA0; margin:6px 0 0; }}
    .item {{ margin:0 0 4px; font-size:15px; line-height:1.65; }}
    .sub-item {{ margin:0 0 4px 22px; font-size:14px; color:#C6CCDA; line-height:1.6; }}
    .gap {{ height:6px; }}
    .act {{ display:flex; gap:10px; margin:0 0 6px; font-size:15px; line-height:1.6; }}
    .act-num {{ color:#F7A01C; font-weight:800; }}
    .no {{ margin:10px 0 0; padding:8px 10px; background:#2A1418; border:1px solid #7F2A32; border-radius:8px;
           color:#FFB4B4; font-size:14px; }}
    nav.toc {{ position:fixed; right:12px; bottom:12px; z-index:20; }}
    nav.toc a {{ display:block; width:34px; height:34px; line-height:34px; text-align:center;
                 background:#1E293B; color:#F7A01C; border-radius:8px; font-size:13px; font-weight:700;
                 margin-top:6px; text-decoration:none; }}
    .dl {{ text-align:center; margin:20px 0; }}
    .dl a {{ display:inline-block; background:#F7A01C; color:#111; font-weight:700; font-size:13px;
             padding:8px 16px; border-radius:8px; text-decoration:none; }}
    footer {{ text-align:center; font-size:11px; color:#5A6072; padding:16px 0 60px; }}
</style>
</head>
<body>
<div class="wrap">
    <header>
        <h1>大轉向資產配置策略</h1>
        <div class="meta">數據基準 2026-09-06 ｜ 台股 9/4 收盤 ｜ 匯率/殖利率 9/5-9/6 ｜ 龍九控股</div>
    </header>

    <div class="dl">
        <a href="grand_pivot_deck.pptx">⬇ 下載簡報檔（pptx，可編輯）</a>
    </div>
{''.join(slides_html)}
    <footer>龍九控股 ｜ 大轉向資產配置策略 v4.1 ｜ 2026-09-06</footer>
</div>
<nav class="toc">
    <a href="#s1">1</a><a href="#s2">2</a><a href="#s3">3</a><a href="#s4">4</a>
    <a href="#s5">5</a><a href="#s6">6</a><a href="#s7">7</a><a href="#s8">8</a>
    <a href="#s9">9</a><a href="#s10">10</a><a href="#s11">11</a>
</nav>
</body>
</html>'''

with open(OUT, 'w', encoding='utf-8') as f:
    f.write(html)
print(f'✅ 已產出 {OUT}（{len(slides_html)} 頁）')

# === 產出後自動發布（2026-10-11 使用者核准）===================================
# 症狀：build_final.py（pptx）與本檔（html）都只產出、不推送 → 週日 17:00 的簡報長期停在
#       「已產出但沒上線」，儀表板「📕 大轉向簡報」按鈕指向前一週版本
#       （10-04 實例：靠當晚晚報的資料 commit 順帶帶上去，不是簡報自己的發布路徑）。
# 修法：最後一棒（本檔）產完後走既有資料通道自發布 —— 只碰本對檔（html ＋ pptx），
#       commit → auto_record（RECORD 通道）→ auto_push（內含 ls-remote 遠端 sha 驗證）。
# 界線：不含任何程式檔；無變更 → 不產生空 commit；發布失敗 → 印 ❌ 且 sys.exit(1)（cron 必須看到紅）。
# 旗標：--no-push 略過發布（本機測試）；--dry-run 只印將執行的動作，不動 git。
# ============================================================================
PUB_FILES = ['grand_pivot_deck.html', 'grand_pivot_deck.pptx']
_PUB_CMD = os.path.join(BASE, 'auto_push.py')


def publish_deck(dry_run=False, no_push=False):
    import subprocess
    import sys as _sys
    if no_push:
        print('🧪 --no-push：已略過發布（本機測試用）')
        return 0
    st = subprocess.run(['git', 'status', '--porcelain', '--', *PUB_FILES], cwd=BASE,
                        capture_output=True, text=True, encoding='utf-8', errors='replace')
    if st.returncode != 0:
        print(f'❌ 發布前檢查失敗（git status rc={st.returncode}）：{(st.stderr or "").strip()[:160]}')
        return 1
    if not st.stdout.strip():
        # 無變更也要確認「遠端已同步」（CIO required_fix #1）：
        #   若上一輪 commit 成功但 push 失敗，本輪會印綠燈卻讓線上停舊版 → 延遲型靜默成功。
        #   先做本地未推數檢查（不依賴網路）；remote 查得到時再比對 sha（查不到僅警告，不製造假紅）。
        _n = subprocess.run(['git', 'rev-list', '--count', 'origin/clean-main..HEAD'], cwd=BASE,
                            capture_output=True, text=True, encoding='utf-8', errors='replace').stdout.strip()
        if _n and _n != '0':
            print(f'❌ 簡報無變更，但本機有 {_n} 顆未推 commit → 線上可能仍是舊版，需處理後再放行')
            return 1
        _head = subprocess.run(['git', 'rev-parse', 'HEAD'], cwd=BASE, capture_output=True, text=True,
                               encoding='utf-8', errors='replace').stdout.strip()
        _ls = subprocess.run(['git', 'ls-remote', 'origin', 'refs/heads/clean-main'], cwd=BASE,
                             capture_output=True, text=True, encoding='utf-8', errors='replace')
        if _ls.returncode == 0 and _ls.stdout.strip():
            _remote = _ls.stdout.split()[0]
            if _remote != _head:
                print(f'❌ 簡報無變更，但遠端 sha {_remote[:12]} ≠ HEAD {_head[:12]} → 未同步，不放行')
                return 1
            print(f'⏭ 簡報無變更，略過發布（遠端已同步 {_remote[:12]}）')
        else:
            print(f'⚠️ 簡報無變更，略過發布（遠端 sha 查詢失敗，本地未推=0 → 僅警告）')
        return 0
    steps = [
        ['git', 'add', '--', *PUB_FILES],
        # pathspec 式 commit（CIO required_fix #2）：只提交這 2 檔，index 若夾帶他班 staged 檔也不會被一起推上去
        ['git', 'commit', '-q', '-m', 'publish(deck): 大轉向簡報自動發布（週日 17:00）', '--', *PUB_FILES],
        [_sys.executable, os.path.join(BASE, 'auto_record.py'), '--script', 'build_final.py',
         '--own', *PUB_FILES, '--note', '週日17:00 簡報自動發布（build_refinance_html 收尾）'],
        [_sys.executable, _PUB_CMD, '--script', 'build_final.py'],
    ]
    if dry_run:
        print('🔍 --dry-run 將執行：')
        for s in steps:
            print('   ' + ' '.join(str(x) for x in s))
        return 0
    for s in steps:
        r = subprocess.run(s, cwd=BASE, capture_output=True, text=True, encoding='utf-8', errors='replace')
        if r.returncode != 0:
            print(f'❌ 簡報發布失敗（step={" ".join(str(x) for x in s[:2])} rc={r.returncode}）')
            tail = ((r.stdout or '') + (r.stderr or '')).strip().splitlines()[-3:]
            for ln in tail:
                print('   ' + ln[:200])
            return 1
        for ln in (r.stdout or '').strip().splitlines():
            if any(k in ln for k in ('✅', '⏭', '⚠️', '❌', 'RECORD')):
                print('   ' + ln.strip()[:200])
    print('✅ 簡報已發布並驗證遠端 sha（誤差 0）')
    return 0


sys.exit(publish_deck(dry_run='--dry-run' in sys.argv, no_push='--no-push' in sys.argv))
